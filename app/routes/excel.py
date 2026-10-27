from __future__ import annotations

import io
import uuid
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from flask import (
    Blueprint, render_template, request, redirect,
    url_for, flash, send_file,
)
from app.db import get_db, get_data_folder, mark_saved
from app.utils import today_iso

excel_bp = Blueprint('excel', __name__)

# ── shared helpers ────────────────────────────────────────────────────────────

def _cell_str(cell) -> str:
    v = cell.value
    return str(v).strip() if v is not None else ''


def _cell_float(cell) -> Optional[float]:
    v = cell.value
    if v is None or str(v).strip() == '':
        return None
    try:
        return float(str(v).strip().replace(',', '').lstrip('₹'))
    except (ValueError, TypeError):
        return None


def _to_paise(val: Optional[float]) -> int:
    if val is None:
        return 0
    return int((Decimal(str(val)) * 100).to_integral_value(rounding=ROUND_HALF_UP))


def _to_grams(val: Optional[float]) -> int:
    if val is None:
        return 0
    return int((Decimal(str(val)) * 1000).to_integral_value(rounding=ROUND_HALF_UP))


def _tmp_dir() -> Path:
    df = get_data_folder()
    d = (df if df else Path('/tmp')) / 'tmp'
    d.mkdir(exist_ok=True)
    return d


def _header_row(ws, headers: list[str]) -> None:
    hdr_font = Font(bold=True, color='FFFFFF')
    hdr_fill = PatternFill('solid', fgColor='2C5F8A')
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal='center')


def _get_or_create_label(db, table: str, name: str) -> Optional[int]:
    if not name:
        return None
    row = db.execute(f"SELECT id FROM {table} WHERE name=?", (name,)).fetchone()
    if row:
        return row['id']
    return db.execute(f"INSERT INTO {table} (name) VALUES (?)", (name,)).lastrowid


def _save_upload(f) -> tuple[str, Path]:
    """Save uploaded file to tmp dir. Returns (tmp_id, path)."""
    tmp_id   = uuid.uuid4().hex[:16]
    tmp_path = _tmp_dir() / f'bm_import_{tmp_id}.xlsx'
    tmp_path.write_bytes(f.stream.read())
    return tmp_id, tmp_path


def _load_tmp(tmp_id: str) -> Optional[Path]:
    p = _tmp_dir() / f'bm_import_{tmp_id}.xlsx'
    return p if p.exists() else None


# ── index ─────────────────────────────────────────────────────────────────────

@excel_bp.route('/excel')
def excel_index():
    return render_template('excel/index.html', today=today_iso())


# ═══════════════════════════════════════════════════════════════════════════════
# PRODUCTS IMPORT / TEMPLATE
# ═══════════════════════════════════════════════════════════════════════════════

_PRODUCT_HEADERS = [
    'SKU', 'Name', 'Material', 'Shape', 'Size (mm)', 'Color',
    'Barcode', 'Cost per kg (₹)', 'Price per kg (₹)',
]

_PRODUCT_EXAMPLES = [
    ['BG-001', 'Glass Round 6mm Blue',   'Glass',   'Round', 6, 'Blue',  '', 250, 550],
    ['BG-002', 'Crystal Oval 8mm Clear', 'Crystal', 'Oval',  8, 'Clear', '', 400, 900],
]

_PRODUCT_NOTES = [
    ('SKU',              'Required. Unique code, e.g. BG-001. Automatically uppercased.'),
    ('Name',             'Required. Full product name.'),
    ('Material',         'Optional. e.g. Glass, Crystal. New values created automatically.'),
    ('Shape',            'Optional. e.g. Round, Oval. New values created automatically.'),
    ('Size (mm)',        'Optional. Numeric size in mm. e.g. 6'),
    ('Color',            'Optional. e.g. Blue, Red. New values created automatically.'),
    ('Barcode',          'Optional. Leave blank to use SKU as barcode.'),
    ('Cost per kg (₹)',  'Optional. Purchase cost per kg. Numbers only.'),
    ('Price per kg (₹)', 'Optional. Selling price per kg. Numbers only.'),
]


@excel_bp.route('/excel/template/products')
def template_products():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Products'
    _header_row(ws, _PRODUCT_HEADERS)
    for i, row in enumerate(_PRODUCT_EXAMPLES, 2):
        for j, val in enumerate(row, 1):
            ws.cell(i, j, val)
    for col, w in enumerate([12, 30, 14, 12, 10, 12, 14, 16, 18], 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = w
    _add_instructions(wb, _PRODUCT_NOTES)
    return _send_wb(wb, 'products_import_template.xlsx')


def _parse_products(db, ws) -> list[dict]:
    rows = []
    for row in ws.iter_rows(min_row=2):
        if all(c.value is None for c in row):
            continue
        sku          = _cell_str(row[0]).upper()
        name         = _cell_str(row[1])
        material_str = _cell_str(row[2])
        shape_str    = _cell_str(row[3])
        size_raw     = _cell_float(row[4])
        color_str    = _cell_str(row[5])
        barcode_str  = _cell_str(row[6])
        cost_raw     = _cell_float(row[7])
        price_raw    = _cell_float(row[8]) if len(row) > 8 else None

        errors = []
        if not sku:   errors.append('SKU required')
        if not name:  errors.append('Name required')

        rows.append({
            'sku': sku, 'name': name,
            'material_str': material_str, 'shape_str': shape_str,
            'size_mm': size_raw, 'color_str': color_str,
            'barcode': barcode_str or sku,
            'cost_paise':  _to_paise(cost_raw),
            'price_paise': _to_paise(price_raw),
            'errors': errors,
        })

    seen: set[str] = set()
    for r in rows:
        if r['sku']:
            if r['sku'] in seen:
                r['errors'].append('Duplicate SKU in file')
            seen.add(r['sku'])

    for r in rows:
        if r['sku'] and db.execute(
            "SELECT 1 FROM products WHERE sku=?", (r['sku'],)
        ).fetchone():
            r['errors'].append('SKU already exists in database')

    return rows


@excel_bp.route('/excel/import/products', methods=['POST'])
def import_products():
    mode   = request.form.get('mode', 'dry_run')
    db     = get_db()
    tmp_id = request.form.get('tmp_id', '').strip()

    if mode == 'commit' and tmp_id:
        p = _load_tmp(tmp_id)
        if not p:
            flash('Import session expired. Please upload again.', 'banner-error')
            return redirect(url_for('excel.excel_index'))
        try:
            wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        except Exception:
            flash('Could not re-open import file.', 'banner-error')
            return redirect(url_for('excel.excel_index'))
    else:
        f = request.files.get('file')
        if not f or not f.filename:
            flash('Please select an Excel (.xlsx) file.', 'banner-error')
            return redirect(url_for('excel.excel_index'))
        tmp_id, p = _save_upload(f)
        try:
            wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        except Exception:
            p.unlink(missing_ok=True)
            flash('Could not read file. Make sure it is a valid .xlsx file.', 'banner-error')
            return redirect(url_for('excel.excel_index'))

    if 'Products' not in wb.sheetnames:
        flash('File must contain a sheet named "Products".', 'banner-error')
        return redirect(url_for('excel.excel_index'))

    rows  = _parse_products(db, wb['Products'])
    if not rows:
        flash('No data rows found.', 'banner-error')
        return redirect(url_for('excel.excel_index'))

    valid   = [r for r in rows if not r['errors']]
    invalid = [r for r in rows if r['errors']]

    if mode == 'commit':
        if not valid:
            flash('No valid rows to import.', 'banner-error')
            return redirect(url_for('excel.excel_index'))
        with db:
            for r in valid:
                mid = _get_or_create_label(db, 'materials', r['material_str'])
                sid = _get_or_create_label(db, 'shapes',    r['shape_str'])
                cid = _get_or_create_label(db, 'colors',    r['color_str'])
                db.execute(
                    """INSERT INTO products
                       (sku, name, material_id, shape_id, size_mm, color_id,
                        barcode, cost_per_kg_paise, price_per_kg_paise)
                       VALUES (?,?,?,?,?,?,?,?,?)""",
                    (r['sku'], r['name'], mid, sid, r['size_mm'], cid,
                     r['barcode'], r['cost_paise'], r['price_paise']),
                )
        mark_saved()
        p.unlink(missing_ok=True)
        n = len(valid)
        msg = f'{n} product{"s" if n != 1 else ""} imported.'
        if invalid:
            msg += f' {len(invalid)} row{"s" if len(invalid) != 1 else ""} skipped.'
        flash(msg, 'banner-success')
        return redirect(url_for('products.product_list'))

    return render_template('excel/preview_products.html',
        rows=rows, valid_count=len(valid), invalid_count=len(invalid), tmp_id=tmp_id,
    )


# ═══════════════════════════════════════════════════════════════════════════════
# PURCHASES IMPORT / TEMPLATE
# ═══════════════════════════════════════════════════════════════════════════════

_PURCHASE_HEADERS = [
    'Challan Number', 'Date (YYYY-MM-DD)', 'Supplier',
    'SKU', 'kg received', 'Cost per kg (₹)',
]

_PURCHASE_EXAMPLES = [
    ['CH-001', '2026-10-26', 'Rajesh Beads', 'BG-001', 10.5, 250],
    ['CH-001', '2026-10-26', 'Rajesh Beads', 'BG-002',  5.0, 400],
    ['CH-002', '2026-10-27', 'Mumbai Traders', 'BG-001', 8.0, 260],
]

_PURCHASE_NOTES = [
    ('Challan Number', 'Required. Rows sharing the same Challan Number become one purchase entry.'),
    ('Date',           'Required. Format: YYYY-MM-DD, e.g. 2026-10-26'),
    ('Supplier',       'Optional. Supplier name. Created automatically if new.'),
    ('SKU',            'Required. Must match an existing product SKU exactly.'),
    ('kg received',    'Required. Quantity received in kg. e.g. 10.5'),
    ('Cost per kg (₹)', 'Optional. Leave blank to use the product\'s existing cost.'),
]


@excel_bp.route('/excel/template/purchases')
def template_purchases():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Purchases'
    _header_row(ws, _PURCHASE_HEADERS)
    for i, row in enumerate(_PURCHASE_EXAMPLES, 2):
        for j, val in enumerate(row, 1):
            ws.cell(i, j, val)
    for col, w in enumerate([16, 16, 18, 12, 12, 16], 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = w
    _add_instructions(wb, _PURCHASE_NOTES)
    return _send_wb(wb, 'purchases_import_template.xlsx')


def _parse_purchases(db, ws) -> list[dict]:
    rows = []
    for row in ws.iter_rows(min_row=2):
        if all(c.value is None for c in row):
            continue
        challan_no   = _cell_str(row[0])
        date_str     = _cell_str(row[1])
        supplier_str = _cell_str(row[2])
        sku          = _cell_str(row[3]).upper()
        kg_raw       = _cell_float(row[4])
        cost_raw     = _cell_float(row[5]) if len(row) > 5 else None

        errors = []
        if not challan_no:  errors.append('Challan number required')
        if not date_str:    errors.append('Date required')
        elif len(date_str) != 10 or date_str[4] != '-':
            errors.append('Date must be YYYY-MM-DD')
        if not sku:         errors.append('SKU required')

        grams = _to_grams(kg_raw)
        if kg_raw is None:  errors.append('kg required')
        elif grams <= 0:    errors.append('kg must be > 0')

        product_id = None
        product_cost = 0
        if sku:
            p = db.execute(
                "SELECT id, cost_per_kg_paise FROM products WHERE sku=? AND active=1", (sku,)
            ).fetchone()
            if p:
                product_id   = p['id']
                product_cost = p['cost_per_kg_paise']
            else:
                errors.append(f'SKU "{sku}" not found')

        cost_paise = _to_paise(cost_raw) if cost_raw is not None else product_cost

        rows.append({
            'challan_no': challan_no, 'date': date_str, 'supplier_str': supplier_str,
            'sku': sku, 'grams': grams, 'cost_paise': cost_paise,
            'product_id': product_id,
            'errors': errors,
        })

    return rows


def _group_purchases(rows: list[dict]) -> list[dict]:
    """Group valid rows by (challan_no, date, supplier_str) → list of challans."""
    groups: dict[tuple, dict] = {}
    for r in rows:
        if r['errors']:
            continue
        key = (r['challan_no'], r['date'], r['supplier_str'])
        if key not in groups:
            groups[key] = {'challan_no': r['challan_no'], 'date': r['date'],
                           'supplier_str': r['supplier_str'], 'lines': []}
        groups[key]['lines'].append(r)
    return list(groups.values())


@excel_bp.route('/excel/import/purchases', methods=['POST'])
def import_purchases():
    mode   = request.form.get('mode', 'dry_run')
    db     = get_db()
    tmp_id = request.form.get('tmp_id', '').strip()

    wb, tmp_id, p, err = _load_wb_for_import(tmp_id, mode, 'Purchases')
    if err:
        flash(err, 'banner-error')
        return redirect(url_for('excel.excel_index'))

    rows   = _parse_purchases(db, wb['Purchases'])
    if not rows:
        flash('No data rows found.', 'banner-error')
        return redirect(url_for('excel.excel_index'))

    valid   = [r for r in rows if not r['errors']]
    invalid = [r for r in rows if r['errors']]
    groups  = _group_purchases(rows)

    if mode == 'commit':
        if not groups:
            flash('No valid rows to import.', 'banner-error')
            return redirect(url_for('excel.excel_index'))
        with db:
            for g in groups:
                supplier_id = None
                if g['supplier_str']:
                    s = db.execute(
                        "SELECT id FROM suppliers WHERE name=?", (g['supplier_str'],)
                    ).fetchone()
                    if s:
                        supplier_id = s['id']
                    else:
                        supplier_id = db.execute(
                            "INSERT INTO suppliers (name) VALUES (?)", (g['supplier_str'],)
                        ).lastrowid

                # Use challan_no directly as po_number (must be unique)
                if db.execute(
                    "SELECT 1 FROM purchase_orders WHERE po_number=?", (g['challan_no'],)
                ).fetchone():
                    # Skip duplicates silently — already imported
                    continue

                po_cur = db.execute(
                    """INSERT INTO purchase_orders (po_number, supplier_id, date, status)
                       VALUES (?,?,?,'Received')""",
                    (g['challan_no'], supplier_id, g['date']),
                )
                po_id = po_cur.lastrowid

                rcpt_cur = db.execute(
                    "INSERT INTO receipts (po_id, date) VALUES (?,?)",
                    (po_id, g['date']),
                )
                receipt_id = rcpt_cur.lastrowid

                for line in g['lines']:
                    po_line_cur = db.execute(
                        """INSERT INTO po_lines
                           (po_id, product_id, grams, cost_per_kg_paise, received_grams)
                           VALUES (?,?,?,?,?)""",
                        (po_id, line['product_id'], line['grams'],
                         line['cost_paise'], line['grams']),
                    )
                    po_line_id = po_line_cur.lastrowid
                    db.execute(
                        """INSERT INTO receipt_lines
                           (receipt_id, po_line_id, product_id, grams, cost_per_kg_paise)
                           VALUES (?,?,?,?,?)""",
                        (receipt_id, po_line_id, line['product_id'],
                         line['grams'], line['cost_paise']),
                    )
                    db.execute(
                        """INSERT INTO stock_movements
                           (product_id, type, grams, ref_type, ref_id, date, note)
                           VALUES (?,?,?,?,?,?,?)""",
                        (line['product_id'], 'purchase_receipt', line['grams'],
                         'receipt', receipt_id, g['date'], g['challan_no']),
                    )
        mark_saved()
        if p: p.unlink(missing_ok=True)
        flash(f'{len(groups)} challan{"s" if len(groups) != 1 else ""} imported ({len(valid)} lines).', 'banner-success')
        return redirect(url_for('po.po_list'))

    return render_template('excel/preview_purchases.html',
        rows=rows, groups=groups, valid_count=len(valid), invalid_count=len(invalid),
        tmp_id=tmp_id,
    )


# ═══════════════════════════════════════════════════════════════════════════════
# SALES IMPORT / TEMPLATE
# ═══════════════════════════════════════════════════════════════════════════════

_SALES_HEADERS = [
    'Invoice Number', 'Date (YYYY-MM-DD)', 'Customer',
    'SKU', 'kg', 'Price per kg (₹)',
]

_SALES_EXAMPLES = [
    ['INV-001', '2026-10-26', 'Meera Traders', 'BG-001', 5.0, 550],
    ['INV-001', '2026-10-26', 'Meera Traders', 'BG-002', 2.5, 900],
    ['INV-002', '2026-10-27', '',              'BG-001', 1.0, 550],
]

_SALES_NOTES = [
    ('Invoice Number', 'Required. Rows sharing the same Invoice Number become one sale.'),
    ('Date',           'Required. Format: YYYY-MM-DD, e.g. 2026-10-26'),
    ('Customer',       'Optional. Leave blank for walk-in. Created automatically if new.'),
    ('SKU',            'Required. Must match an existing product SKU exactly.'),
    ('kg',             'Required. Quantity sold in kg.'),
    ('Price per kg (₹)', 'Optional. Leave blank to use the product\'s selling price.'),
]


@excel_bp.route('/excel/template/sales')
def template_sales():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sales'
    _header_row(ws, _SALES_HEADERS)
    for i, row in enumerate(_SALES_EXAMPLES, 2):
        for j, val in enumerate(row, 1):
            ws.cell(i, j, val)
    for col, w in enumerate([16, 16, 18, 12, 8, 16], 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = w
    _add_instructions(wb, _SALES_NOTES)
    return _send_wb(wb, 'sales_import_template.xlsx')


def _parse_sales(db, ws) -> list[dict]:
    rows = []
    for row in ws.iter_rows(min_row=2):
        if all(c.value is None for c in row):
            continue
        invoice_no   = _cell_str(row[0])
        date_str     = _cell_str(row[1])
        customer_str = _cell_str(row[2])
        sku          = _cell_str(row[3]).upper()
        kg_raw       = _cell_float(row[4])
        price_raw    = _cell_float(row[5]) if len(row) > 5 else None

        errors = []
        if not invoice_no:  errors.append('Invoice number required')
        if not date_str:    errors.append('Date required')
        elif len(date_str) != 10 or date_str[4] != '-':
            errors.append('Date must be YYYY-MM-DD')
        if not sku:         errors.append('SKU required')

        grams = _to_grams(kg_raw)
        if kg_raw is None:  errors.append('kg required')
        elif grams <= 0:    errors.append('kg must be > 0')

        product_id    = None
        product_price = 0
        if sku:
            p = db.execute(
                "SELECT id, price_per_kg_paise FROM products WHERE sku=? AND active=1", (sku,)
            ).fetchone()
            if p:
                product_id    = p['id']
                product_price = p['price_per_kg_paise']
            else:
                errors.append(f'SKU "{sku}" not found')

        price_paise = _to_paise(price_raw) if price_raw is not None else product_price

        rows.append({
            'invoice_no': invoice_no, 'date': date_str, 'customer_str': customer_str,
            'sku': sku, 'grams': grams, 'price_paise': price_paise,
            'product_id': product_id,
            'errors': errors,
        })

    # Flag invoice numbers that already exist in DB
    seen_invoices: set[str] = set()
    for r in rows:
        if not r['invoice_no']:
            continue
        if r['invoice_no'] in seen_invoices:
            continue
        seen_invoices.add(r['invoice_no'])
        if db.execute(
            "SELECT 1 FROM sales WHERE invoice_number=?", (r['invoice_no'],)
        ).fetchone():
            for row2 in rows:
                if row2['invoice_no'] == r['invoice_no']:
                    row2['errors'].append(f'Invoice "{r["invoice_no"]}" already exists')

    return rows


def _group_sales(rows: list[dict]) -> list[dict]:
    groups: dict[tuple, dict] = {}
    for r in rows:
        if r['errors']:
            continue
        key = (r['invoice_no'], r['date'], r['customer_str'])
        if key not in groups:
            groups[key] = {'invoice_no': r['invoice_no'], 'date': r['date'],
                           'customer_str': r['customer_str'], 'lines': []}
        groups[key]['lines'].append(r)
    return list(groups.values())


@excel_bp.route('/excel/import/sales', methods=['POST'])
def import_sales():
    mode   = request.form.get('mode', 'dry_run')
    db     = get_db()
    tmp_id = request.form.get('tmp_id', '').strip()

    wb, tmp_id, p, err = _load_wb_for_import(tmp_id, mode, 'Sales')
    if err:
        flash(err, 'banner-error')
        return redirect(url_for('excel.excel_index'))

    rows   = _parse_sales(db, wb['Sales'])
    if not rows:
        flash('No data rows found.', 'banner-error')
        return redirect(url_for('excel.excel_index'))

    valid   = [r for r in rows if not r['errors']]
    invalid = [r for r in rows if r['errors']]
    groups  = _group_sales(rows)

    if mode == 'commit':
        if not groups:
            flash('No valid rows to import.', 'banner-error')
            return redirect(url_for('excel.excel_index'))
        with db:
            for g in groups:
                customer_id = None
                if g['customer_str']:
                    c = db.execute(
                        "SELECT id FROM customers WHERE name=?", (g['customer_str'],)
                    ).fetchone()
                    customer_id = c['id'] if c else db.execute(
                        "INSERT INTO customers (name) VALUES (?)", (g['customer_str'],)
                    ).lastrowid

                sale_cur = db.execute(
                    "INSERT INTO sales (invoice_number, customer_id, date) VALUES (?,?,?)",
                    (g['invoice_no'], customer_id, g['date']),
                )
                sale_id = sale_cur.lastrowid

                for line in g['lines']:
                    db.execute(
                        """INSERT INTO sale_lines
                           (sale_id, product_id, grams, price_per_kg_paise)
                           VALUES (?,?,?,?)""",
                        (sale_id, line['product_id'], line['grams'], line['price_paise']),
                    )
                    db.execute(
                        """INSERT INTO stock_movements
                           (product_id, type, grams, ref_type, ref_id, date)
                           VALUES (?,?,?,?,?,?)""",
                        (line['product_id'], 'sale', -line['grams'],
                         'sale', sale_id, g['date']),
                    )
        mark_saved()
        if p: p.unlink(missing_ok=True)
        flash(f'{len(groups)} sale{"s" if len(groups) != 1 else ""} imported ({len(valid)} lines).', 'banner-success')
        return redirect(url_for('sales.sale_list'))

    return render_template('excel/preview_sales.html',
        rows=rows, groups=groups, valid_count=len(valid), invalid_count=len(invalid),
        tmp_id=tmp_id,
    )


# ═══════════════════════════════════════════════════════════════════════════════
# EXPORTS
# ═══════════════════════════════════════════════════════════════════════════════

@excel_bp.route('/excel/export/products')
def export_products():
    db   = get_db()
    rows = db.execute("""
        SELECT p.sku, p.name,
               m.name AS material, s.name AS shape,
               p.size_mm, c.name AS color, p.barcode,
               p.cost_per_kg_paise, p.price_per_kg_paise,
               COALESCE(SUM(sm.grams), 0) AS stock_grams,
               p.reorder_grams, p.active
        FROM products p
        LEFT JOIN materials m ON p.material_id = m.id
        LEFT JOIN shapes    s ON p.shape_id    = s.id
        LEFT JOIN colors    c ON p.color_id    = c.id
        LEFT JOIN stock_movements sm ON sm.product_id = p.id
        GROUP BY p.id
        ORDER BY p.sku
    """).fetchall()

    wb  = openpyxl.Workbook()
    ws  = wb.active
    ws.title = 'Products'
    _header_row(ws, [
        'SKU', 'Name', 'Material', 'Shape', 'Size (mm)', 'Color',
        'Cost per kg (₹)', 'Price per kg (₹)', 'Stock (kg)',
        'Reorder at (kg)', 'Barcode', 'Active',
    ])
    for i, r in enumerate(rows, 2):
        ws.cell(i, 1,  r['sku'])
        ws.cell(i, 2,  r['name'])
        ws.cell(i, 3,  r['material'] or '')
        ws.cell(i, 4,  r['shape']    or '')
        ws.cell(i, 5,  r['size_mm'])
        ws.cell(i, 6,  r['color']    or '')
        ws.cell(i, 7,  round(r['cost_per_kg_paise']  / 100, 2))
        ws.cell(i, 8,  round(r['price_per_kg_paise'] / 100, 2))
        ws.cell(i, 9,  round(r['stock_grams'] / 1000, 3))
        ws.cell(i, 10, round(r['reorder_grams'] / 1000, 3))
        ws.cell(i, 11, r['barcode'] or '')
        ws.cell(i, 12, 'Yes' if r['active'] else 'No')
    return _send_wb(wb, f'products_{today_iso()}.xlsx')


@excel_bp.route('/excel/export/sales')
def export_sales():
    db        = get_db()
    date_from = request.args.get('from', '')
    date_to   = request.args.get('to',   today_iso())

    where  = ['s.voided = 0']
    params: list = []
    if date_from:
        where.append('s.date >= ?'); params.append(date_from)
    if date_to:
        where.append('s.date <= ?'); params.append(date_to)

    rows = db.execute(f"""
        SELECT s.invoice_number, s.date,
               COALESCE(c.name, '— Walk-in —') AS customer,
               COALESCE(SUM(sl.grams), 0) AS total_grams,
               COALESCE(SUM((sl.price_per_kg_paise * sl.grams + 500) / 1000), 0) AS total_paise,
               COALESCE(pt.paid, 0) AS paid_paise
        FROM sales s
        LEFT JOIN customers c    ON s.customer_id = c.id
        LEFT JOIN sale_lines sl  ON sl.sale_id    = s.id
        LEFT JOIN (
            SELECT sale_id, SUM(paise) AS paid
            FROM payments WHERE voided=0 GROUP BY sale_id
        ) pt ON pt.sale_id = s.id
        WHERE {' AND '.join(where)}
        GROUP BY s.id
        ORDER BY s.date, s.invoice_number
    """, params).fetchall()

    wb  = openpyxl.Workbook()
    ws  = wb.active
    ws.title = 'Sales'
    _header_row(ws, ['Invoice', 'Date', 'Customer', 'kg', 'Total (₹)', 'Paid (₹)', 'Balance (₹)', 'Status'])
    for i, r in enumerate(rows, 2):
        total_p = r['total_paise']
        paid_p  = r['paid_paise']
        balance = max(0, total_p - paid_p)
        if paid_p >= total_p: status = 'Paid'
        elif paid_p > 0:      status = 'Part paid'
        else:                 status = 'Pending'
        ws.cell(i, 1, r['invoice_number'])
        ws.cell(i, 2, r['date'])
        ws.cell(i, 3, r['customer'])
        ws.cell(i, 4, round(r['total_grams'] / 1000, 3))
        ws.cell(i, 5, round(total_p / 100, 2))
        ws.cell(i, 6, round(paid_p  / 100, 2))
        ws.cell(i, 7, round(balance / 100, 2))
        ws.cell(i, 8, status)
    return _send_wb(wb, f'sales_{date_from or "all"}_{date_to}.xlsx')


@excel_bp.route('/excel/export/stock')
def export_stock():
    db   = get_db()
    rows = db.execute("""
        SELECT p.sku, p.name,
               m.name AS material, c.name AS color,
               COALESCE(SUM(sm.grams), 0) AS stock_grams,
               p.cost_per_kg_paise,
               COALESCE(SUM(CAST(sm.grams AS REAL) * p.cost_per_kg_paise / 1000.0), 0) AS value_paise
        FROM products p
        LEFT JOIN materials m ON p.material_id = m.id
        LEFT JOIN colors    c ON p.color_id    = c.id
        LEFT JOIN stock_movements sm ON sm.product_id = p.id
        WHERE p.active = 1
        GROUP BY p.id
        ORDER BY p.sku
    """).fetchall()

    wb  = openpyxl.Workbook()
    ws  = wb.active
    ws.title = 'Stock Value'
    _header_row(ws, ['SKU', 'Name', 'Material', 'Color', 'Stock (kg)', 'Cost per kg (₹)', 'Value at cost (₹)'])
    for i, r in enumerate(rows, 2):
        ws.cell(i, 1, r['sku'])
        ws.cell(i, 2, r['name'])
        ws.cell(i, 3, r['material'] or '')
        ws.cell(i, 4, r['color']    or '')
        ws.cell(i, 5, round(r['stock_grams'] / 1000, 3))
        ws.cell(i, 6, round(r['cost_per_kg_paise'] / 100, 2))
        ws.cell(i, 7, round(r['value_paise'] / 100, 2))
    return _send_wb(wb, f'stock_value_{today_iso()}.xlsx')


# ═══════════════════════════════════════════════════════════════════════════════
# SHARED UTILITIES
# ═══════════════════════════════════════════════════════════════════════════════

def _add_instructions(wb, notes: list[tuple[str, str]]) -> None:
    inst = wb.create_sheet('Instructions')
    inst.cell(1, 1, 'Column').font    = Font(bold=True)
    inst.cell(1, 2, 'Description').font = Font(bold=True)
    for i, (col_name, desc) in enumerate(notes, 2):
        inst.cell(i, 1, col_name)
        inst.cell(i, 2, desc)
    inst.column_dimensions['A'].width = 22
    inst.column_dimensions['B'].width = 65


def _send_wb(wb, filename: str):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf,
        download_name=filename,
        as_attachment=True,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def _load_wb_for_import(tmp_id: str, mode: str, sheet_name: str):
    """
    Returns (wb, tmp_id, path, error_str).
    error_str is non-None when something went wrong.
    """
    if mode == 'commit' and tmp_id:
        p = _load_tmp(tmp_id)
        if not p:
            return None, tmp_id, None, 'Import session expired. Please upload again.'
        try:
            wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        except Exception:
            return None, tmp_id, None, 'Could not re-open import file.'
    else:
        f = request.files.get('file')
        if not f or not f.filename:
            return None, '', None, 'Please select an Excel (.xlsx) file.'
        tmp_id, p = _save_upload(f)
        try:
            wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        except Exception:
            p.unlink(missing_ok=True)
            return None, '', None, 'Could not read file. Make sure it is a valid .xlsx file.'

    if sheet_name not in wb.sheetnames:
        return None, tmp_id, p, f'File must contain a sheet named "{sheet_name}".'

    return wb, tmp_id, p, None
